from io import BytesIO
import re

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from starlette.responses import StreamingResponse

MAX_EXCEL_SIZE = 10 * 1024 * 1024


def read_excel_rows(contents: bytes) -> list[tuple[int, dict[str, object]]]:
    if not contents or len(contents) > MAX_EXCEL_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail="Excel files must be non-empty and no larger than 10 MB.",
        )

    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers_row = next(values, None)
        if not headers_row:
            raise ValueError("The worksheet is empty.")

        headers = [
            re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
            for value in headers_row
        ]
        if not any(headers):
            raise ValueError("The first row must contain column headers.")

        rows = []
        for row_number, row in enumerate(values, start=2):
            record = {
                header: value
                for header, value in zip(headers, row)
                if header and value is not None
            }
            if any(value is not None and str(value).strip() for value in row):
                rows.append((row_number, record))
        workbook.close()
        return rows
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unable to read this workbook. Upload a valid .xlsx file.",
        ) from exc


def excel_template(
    filename: str,
    headers: list[str],
    sample: list[object] | None = None,
):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Import"
    worksheet.append(headers)
    if sample is not None:
        worksheet.append(sample)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column in worksheet.columns:
        cells = list(column)
        width = max(14, min(32, max(len(str(cell.value or "")) for cell in cells) + 2))
        worksheet.column_dimensions[cells[0].column_letter].width = width

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
